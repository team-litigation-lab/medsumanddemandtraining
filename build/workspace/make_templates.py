#!/usr/bin/env python3
"""The Google Workspace templates for the Dana Whitfield case workspace.

    PYTHONPATH=<dir with openpyxl> python3 build/workspace/make_templates.py

Writes build/workspace/templates/:
  start_here.html      → "00 START HERE — Your tasks, Days 1–5" (Google Doc)
  file_audit.html      → "Whitfield — Day 1 File Audit" (Google Doc)
  medsum.html          → "Whitfield — Medical Summary" (Google Doc)
  demand_letter.html   → "Whitfield — Demand Letter (draft)" (Google Doc)
  exhibit_index.html   → "Whitfield — Exhibit Index" (Google Doc)
  reply_keystone.html  → "Whitfield — Reply to Keystone (draft)" (Google Doc)
  chronology.xlsx      → "Whitfield — Medical Chronology" (Google Sheet)
  itemization.xlsx     → "Whitfield — Bills Itemization" (Google Sheet)
HTML converts to a Google Doc and .xlsx to a Google Sheet when uploaded to Drive. The templates give
structure only: what to write and where; the answers come from the records.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "templates")
NAVY, ORANGE, FILL = "262B45", "DB8437", "FFF2CC"
PH = 'style="background:#FFF2CC"'           # a placeholder to replace
NOTE = 'style="color:#5B6178;font-style:italic"'


def doc(title, body):
    return f"""<!DOCTYPE html><html><head><meta charset="UTF-8"><title>{title}</title></head>
<body style="font-family:Arial,sans-serif;font-size:11pt;line-height:1.4">{body}</body></html>"""


def table(headers, rows=8, widths=None):
    th = "".join(f'<th style="background:#262B45;color:#ffffff;border:1px solid #999;padding:4px;text-align:left">{h}</th>' for h in headers)
    tr = "".join("<tr>" + "".join('<td style="border:1px solid #999;padding:4px">&nbsp;</td>' for _ in headers) + "</tr>" for _ in range(rows))
    return f'<table style="border-collapse:collapse;width:100%"><tr>{th}</tr>{tr}</table>'


def start_here():
    day = lambda n, t, items, out: (f'<h2 style="color:#262B45">Day {n} — {t}</h2><ol>' + "".join(f"<li>{x}</li>" for x in items) + "</ol>"
                                     f'<p><b>Hand in:</b> {out}</p>')
    return doc("START HERE", f"""
<p style="color:#B54A3F;font-weight:bold">TRAINING — SIMULATED CASE · fictional people and records</p>
<h1 style="color:#262B45">Dana Whitfield v. Keystone Mutual — your case workspace</h1>
<p>This folder is your own copy of a real-looking file, exactly as it reaches a Demand Specialist: one PDF per provider or
sender, unsorted, with fax and scanner names and no Bates numbers yet. Over the five days you'll organize it and build the
work product, from the chronology to the demand letter, in Google Docs and Sheets.</p>
<p><b>Your trainer can open this folder at any time</b> and see how you arranged it, and each document's version history shows
how you built it. Work here, not in a copy on your own computer. When a day's work is ready, press <b>Submit</b> on the
Case Workspace page in the training portal.</p>
<h2 style="color:#262B45">How the folder works</h2>
<ul>
<li><b>01 Incoming — unsorted</b>: everything that has come in so far. Empty it on Day 1.</li>
<li><b>02 Case File</b>: where every document belongs, one subfolder per category.</li>
<li><b>03 Work Product</b>: the documents you write (templates are already there).</li>
<li><b>04 Received after the demand</b>: don't open it until Day 5.</li>
</ul>
<h2 style="color:#262B45">Naming convention (LSH)</h2>
<ul>
<li>Every file: <b>YYYY-MM-DD Source – Document</b>, using the document's own date. Example: <i>2026-04-02 Keystone Mutual – Liability acceptance.pdf</i></li>
<li>Medical records, once you've put them in order: <b>WHITFIELD 0001-0009 Riverside Medical Center – ED record 2026-03-14.pdf</b>
(the Bates range first, so the folder sorts in Bates order).</li>
<li>Never delete a file. Anything you won't use goes to <b>02 Case File / 09 Duplicates &amp; not used</b>, with the reason in your audit.</li>
</ul>
{day(1, "Organize the file", [
    "Open every file in <b>01 Incoming</b> and work out what it is, who sent it and its date.",
    "Rename it with the LSH convention and move it to the right <b>02 Case File</b> subfolder.",
    "Put the medical records in Bates order: by provider, in the order Dana first saw each one after the crash, each provider's pages in order. "
    "Number them <b>WHITFIELD 0001</b> onward and put each file's range at the front of its name.",
    "Check every core fact against the documents (date of incident, names, claim number) and record every problem you find in "
    "<b>Whitfield — Day 1 File Audit</b>: the issue, where it is (file and page) and what you'll do or who you'll tell."],
    "the organized folder and the Day 1 File Audit.")}
{day(2, "Medical chronology and medical summary", [
    "Build <b>Whitfield — Medical Chronology</b>: one row per encounter, in date order, in the provider's words, with the Bates page for every row and a flag where one applies.",
    "Write <b>Whitfield — Medical Summary</b> from the chronology: neutral, factual, a Bates cite for every statement."],
    "the chronology and the medical summary.")}
{day(3, "Bills itemization", [
    "Build <b>Whitfield — Bills Itemization</b> from the bills and ledgers: one line per charge, billed, adjustments, every payment by payer, balance and source.",
    "List what you left out (and why) on the <b>Exclusions</b> tab, and every balance or lien on the <b>Balances &amp; Liens</b> tab."],
    "the itemization (the totals must tie out).")}
{day(4, "The demand letter", [
    "Attorney Bennett's instructions (email of 10/08/2026): <i>“Draft the Whitfield demand to Tom Reyes at Keystone. Use $85,000.00 and give "
    "them 30 days from the date it's received to respond. Address the January 2025 low back strain and the gap in treatment head-on. "
    "I'll review and sign.”</i>",
    "Draft <b>Whitfield — Demand Letter (draft)</b> on the LSH structure. Every fact from the records gets an exhibit and Bates cite; "
    "every figure ties to your itemization; no opinions about value or the law."],
    "the demand letter draft.")}
{day(5, "The packet and the response", [
    "Build <b>Whitfield — Exhibit Index</b> in the LSH standard order, with Bates ranges and page counts.",
    "In <b>02 Case File / 07 Demand Packet</b>, add a shortcut to each exhibit in packet order (right-click → Organize → Add shortcut), named <b>Ex. A – …</b>, <b>Ex. B – …</b>.",
    "Now open <b>04 Received after the demand</b>. File each letter in <b>08 Correspondence</b> with the naming convention.",
    "Draft <b>Whitfield — Reply to Keystone (draft)</b>: answer each point of Keystone's 11/04/2026 letter with the record that answers it. "
    "The attorney decides any counteroffer and the scope of the records request: flag those for her, don't decide them."],
    "the exhibit index, the packet folder and the reply draft.")}
<h2 style="color:#262B45">Rules that never change</h2>
<ul><li>Cite the record for every fact: exhibit and Bates page.</li>
<li>Use the provider's words. Don't upgrade a diagnosis (a protrusion is not a herniation).</li>
<li>Value, the demand amount, offers, counteroffers and legal questions belong to Attorney Bennett. Liens and disbursement belong to Marcus Webb.</li>
<li>Client information stays in this folder. Don't download it to a personal device or share the folder.</li></ul>
""")


def file_audit():
    return doc("Day 1 File Audit", f"""
<p style="color:#B54A3F;font-weight:bold">TRAINING — SIMULATED CASE</p>
<h1 style="color:#262B45">Whitfield — Day 1 File Audit</h1>
<p {NOTE}>Everything you checked or found while organizing the file. One row per problem: what's wrong, where it is, and what you'll do
about it (request, flag for the attorney, tell the Case Manager, correct the intake…).</p>
<p><b>Prepared by:</b> <span {PH}>[your name]</span> · <b>Date:</b> <span {PH}>[date]</span></p>
<h2 style="color:#262B45">Core facts checked</h2>
{table(["Fact", "Intake says", "Documents say", "Source (file, page)", "OK?"], 6)}
<h2 style="color:#262B45">Problems found</h2>
{table(["#", "Problem", "Where (file, page)", "Action / who you'll tell"], 10)}
<h2 style="color:#262B45">Records still missing or incomplete</h2>
{table(["Provider", "What's missing", "How you'll request it"], 4)}
""")


def medsum():
    sec = lambda h, guide: f'<h2 style="color:#262B45">{h}</h2><p {NOTE}>{guide}</p><p {PH}>[write here]</p>'
    return doc("Medical Summary", f"""
<p style="color:#B54A3F;font-weight:bold">TRAINING — SIMULATED CASE · DRAFT — for attorney review</p>
<h1 style="color:#262B45">Medical Summary — Dana Whitfield</h1>
<p><b>Date of incident:</b> <span {PH}>[ ]</span> · <b>Records reviewed:</b> <span {PH}>[WHITFIELD ____–____]</span> · <b>Prepared by:</b> <span {PH}>[ ]</span></p>
{sec("Overview", "Who the client is, the date of incident, how it happened, the treatment span, the providers and the number of visits.")}
{sec("Initial treatment", "The emergency department or first visit: complaints, tests, diagnoses. Cite the pages.")}
{sec("Diagnostics", "Imaging, in the radiologist's own words.")}
{sec("Treatment course", "By provider or by phase: what was done and how the client responded.")}
{sec("Prior history and gaps in treatment", "Stated plainly, with the reasons the records give.")}
{sec("Current status and future care", "MMI, restrictions, the future care recommended and its cost, from the provider's note.")}
<p {NOTE}>Neutral and factual. A Bates cite for every statement.</p>
""")


def demand_letter():
    sec = lambda h, guide: f'<h3 style="color:#262B45">{h}</h3><p {NOTE}>{guide}</p><p {PH}>[write here]</p>'
    return doc("Demand Letter", f"""
<p style="color:#B54A3F;font-weight:bold">TRAINING — SIMULATED CASE · DRAFT — NOT FOR RELEASE · for attorney review</p>
<p style="font-size:16pt;font-weight:bold;color:#262B45;margin-bottom:0">LSH LAW GROUP</p>
<p style="color:#5B6178;margin-top:0">500 Lakeview Plaza, Suite 1200 · Lakeside, ST 90310 · (555) 318-1000</p>
<p><span {PH}>[date]</span></p>
<p><span {PH}>[delivery method]</span><br><span {PH}>[adjuster name, title]</span><br><span {PH}>[carrier]</span><br><span {PH}>[address]</span></p>
<p><b>RE:</b> Our client: <span {PH}>[ ]</span><br>Your insured: <span {PH}>[ ]</span><br>Claim No.: <span {PH}>[ ]</span><br>Date of loss: <span {PH}>[ ]</span></p>
<p>Dear <span {PH}>[ ]</span>:</p>
{sec("Introduction", "Who we represent and the purpose of the letter.")}
{sec("Facts and liability", "What happened and why the insured is responsible. Cite the police report and the photos (exhibits).")}
{sec("Injuries and treatment", "From your medical summary: the injuries and the treatment in date order, with exhibit and Bates cites.")}
<h3 style="color:#262B45">Medical specials</h3><p {NOTE}>From your itemization, to the cent. Name the exhibit.</p>
{table(["Provider", "Dates of service", "Amount"], 9)}
<p><b>Total past medical expenses:</b> <span {PH}>[ ]</span> (Ex. <span {PH}>[ ]</span>)</p>
{sec("Future medical care", "Only what a provider recommended, with the cost they gave and the cite.")}
{sec("Lost wages", "The days, the rate and the total, with the employer's verification as the exhibit.")}
{sec("Pain, suffering and loss of enjoyment of life", "Specific, truthful examples from the client's signed impact statement and the records.")}
{sec("The records, addressed", "Any prior history and any gap in treatment, stated plainly with the record that explains it.")}
{sec("Demand", "The amount and the response deadline the attorney set.")}
<p>Sincerely,</p><p><br>Laura Bennett, Esq.<br>LSH Law Group</p>
<p><b>Enclosures:</b> <span {PH}>[exhibit list]</span></p>
""")


def exhibit_index():
    return doc("Exhibit Index", f"""
<p style="color:#B54A3F;font-weight:bold">TRAINING — SIMULATED CASE</p>
<h1 style="color:#262B45">Exhibit Index — Whitfield Demand</h1>
<p {NOTE}>In the LSH standard order. Bates ranges and page counts must match the PDFs exactly.</p>
{table(["Exhibit", "Description", "Bates range / pages", "File in 02 Case File"], 9)}
<h2 style="color:#262B45">Packet checks</h2>
<ul><li>☐ Every cite in the letter lands on the right page.</li><li>☐ Figures match the itemization to the cent.</li>
<li>☐ Exhibits in order; shortcuts in 07 Demand Packet named Ex. A, Ex. B…</li><li>☐ Nothing in the packet that doesn't belong (duplicates, unrelated records).</li></ul>
""")


def reply_keystone():
    return doc("Reply to Keystone", f"""
<p style="color:#B54A3F;font-weight:bold">TRAINING — SIMULATED CASE · DRAFT — for attorney review</p>
<h1 style="color:#262B45">Reply to Keystone — draft</h1>
<p><b>To:</b> <span {PH}>[ ]</span> · <b>Re:</b> <span {PH}>[ ]</span> · <b>Their letter dated:</b> <span {PH}>[ ]</span></p>
<h2 style="color:#262B45">Each point they raised, and the record that answers it</h2>
{table(["Their argument (their words)", "What the records show", "Exhibit / Bates"], 6)}
<h2 style="color:#262B45">Draft reply paragraphs</h2><p {PH}>[write here]</p>
<h2 style="color:#262B45">For Attorney Bennett to decide</h2>
<p {NOTE}>The offer, any counteroffer, and the scope of the records request. List what she needs to decide and the facts she needs; don't decide them.</p>
<p {PH}>[write here]</p>
""")


def sheet_header(ws, headers, widths):
    thin = Side(style="thin", color="999999")
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30


def chronology():
    wb = Workbook(); ws = wb.active; ws.title = "Chronology"
    heads = ["Date of service", "Provider / facility", "Visit type", "Complaints & findings (provider's words)", "Diagnosis", "Treatment / plan", "Flag", "Bates page(s)"]
    sheet_header(ws, heads, [14, 26, 18, 52, 30, 34, 22, 16])
    flags = DataValidation(type="list", formula1='"Routine entry,Prior injury,Gap in treatment,Objective finding,Causation opinion,Procedure,MMI / future care"', allow_blank=True)
    ws.add_data_validation(flags); flags.add("G2:G200")
    for r in range(2, 201):
        for col in "ADEF": ws[f"{col}{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        ws[f"A{r}"].number_format = "mm/dd/yyyy"
    n = wb.create_sheet("Notes")
    for i, t in enumerate(["One row per encounter, in date order (prior records first).", "Use the provider's words. Cite the Bates page for every row.",
                           "Flag: Prior injury · Gap in treatment (30+ days) · Objective finding · Causation opinion · Procedure · MMI / future care."], 1):
        n.cell(row=i, column=1, value=t)
    n.column_dimensions["A"].width = 110
    wb.save(os.path.join(OUT, "chronology.xlsx"))


def itemization():
    wb = Workbook(); ws = wb.active; ws.title = "Itemization"
    heads = ["Provider", "Date(s) of service", "Code (CPT / rev.)", "Description", "Billed", "Adjustment", "Health ins. paid", "PIP / MedPay paid", "Client paid", "Balance", "Source (file, page)", "Notes"]
    sheet_header(ws, heads, [28, 18, 14, 36, 13, 13, 14, 14, 12, 13, 26, 30])
    money = '"$"#,##0.00'
    last = 41
    for r in range(2, last):
        ws[f"J{r}"] = f'=IF(COUNT(E{r}:I{r})=0,"",E{r}-F{r}-G{r}-H{r}-I{r})'
        for col in "EFGHIJ": ws[f"{col}{r}"].number_format = money
    t = last
    ws[f"A{t}"] = "TOTALS"; ws[f"A{t}"].font = Font(bold=True)
    for col in "EFGHIJ":
        ws[f"{col}{t}"] = f"=SUM({col}2:{col}{t - 1})"; ws[f"{col}{t}"].number_format = money
        ws[f"{col}{t}"].font = Font(bold=True); ws[f"{col}{t}"].fill = PatternFill("solid", fgColor="FFF6EC")
    ws[f"A{t + 1}"] = "Check: Billed − Adjustments − Payments = Balance →"
    ws[f"J{t + 1}"] = f'=IF(ROUND(E{t}-F{t}-G{t}-H{t}-I{t}-J{t},2)=0,"ties out","does NOT tie out")'
    ex = wb.create_sheet("Exclusions")
    sheet_header(ex, ["Provider", "Date of service", "Description", "Amount", "Why it's left out", "Source (file, page)"], [28, 16, 36, 13, 44, 26])
    for r in range(2, 20): ex[f"D{r}"].number_format = money
    li = wb.create_sheet("Balances & Liens")
    sheet_header(li, ["Holder", "Type (LOP, balance due, reimbursement claim)", "Amount", "Confirmed? (date, how)", "Source (file, page)"], [30, 36, 14, 26, 26])
    for r in range(2, 20): li[f"C{r}"].number_format = money
    wb.save(os.path.join(OUT, "itemization.xlsx"))


def main():
    os.makedirs(OUT, exist_ok=True)
    for name, fn in [("start_here", start_here), ("file_audit", file_audit), ("medsum", medsum), ("demand_letter", demand_letter),
                     ("exhibit_index", exhibit_index), ("reply_keystone", reply_keystone)]:
        open(os.path.join(OUT, f"{name}.html"), "w", encoding="utf8").write(fn())
    chronology(); itemization()
    for f in sorted(os.listdir(OUT)): print(f, os.path.getsize(os.path.join(OUT, f)), "bytes")


if __name__ == "__main__":
    main()
